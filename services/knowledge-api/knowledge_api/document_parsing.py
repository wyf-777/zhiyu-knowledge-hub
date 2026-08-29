from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from importlib.util import find_spec
from pathlib import Path
from typing import Any, Callable, Literal, Protocol

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation

TARGET_CHUNK_CHARS = 900
CHUNK_OVERLAP_CHARS = 120
PARENT_CHUNK_CHARS = 3600
MIN_PDF_TEXT_CHARS = 20


@dataclass
class TextSection:
    text: str
    page: int | None = None
    heading: str = ""
    content_kind: Literal["text", "table", "ocr"] = "text"
    location: dict[str, Any] | None = None


@dataclass(frozen=True)
class OcrStatus:
    provider: str
    state: Literal["ready", "unavailable", "error"]
    execution_mode: str
    message: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


class OcrUnavailableError(RuntimeError):
    pass


class OcrEngine(Protocol):
    def status(self) -> OcrStatus: ...

    def recognize_pdf_pages(self, path: Path, page_numbers: list[int]) -> dict[int, str]: ...


class RapidOcrEngine:
    """Lazy, fully local OCR adapter shared by Windows and macOS Sidecars."""

    def __init__(self) -> None:
        self._engine: Any | None = None
        self._last_error = ""

    def status(self) -> OcrStatus:
        available = find_spec("pymupdf") is not None and find_spec("rapidocr_onnxruntime") is not None
        if not available:
            return OcrStatus("rapidocr", "unavailable", "local-cpu", "本地 OCR 组件未安装，扫描 PDF 暂不可解析。")
        if self._last_error:
            return OcrStatus("rapidocr", "error", "local-cpu", f"本地 OCR 最近一次运行失败：{self._last_error}")
        return OcrStatus("rapidocr", "ready", "local-cpu", "扫描 PDF 将使用本机 RapidOCR，不上传云端。")

    def _get_engine(self) -> Any:
        if self.status().state == "unavailable":
            raise OcrUnavailableError("本地 OCR 组件不可用，请重新安装知识库服务依赖")
        if self._engine is None:
            from rapidocr_onnxruntime import RapidOCR

            self._engine = RapidOCR()
        return self._engine

    def recognize_pdf_pages(self, path: Path, page_numbers: list[int]) -> dict[int, str]:
        try:
            import pymupdf

            engine = self._get_engine()
            output: dict[int, str] = {}
            with pymupdf.open(path) as document:
                for page_number in page_numbers:
                    page = document.load_page(page_number - 1)
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.0, 2.0), colorspace=pymupdf.csRGB, alpha=False)
                    result, _ = engine(pixmap.tobytes("png"))
                    lines = [
                        str(item[1]).strip()
                        for item in (result or [])
                        if len(item) >= 3 and str(item[1]).strip() and float(item[2]) >= 0.45
                    ]
                    output[page_number] = "\n".join(lines)
            self._last_error = ""
            return output
        except OcrUnavailableError:
            raise
        except Exception as error:  # noqa: BLE001
            self._last_error = str(error)[:240]
            raise RuntimeError(f"本地 OCR 处理失败：{self._last_error}") from error


def _clean_cell(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def _render_table(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    if len(normalized) == 1:
        normalized.insert(0, [f"字段{index + 1}" for index in range(width)])

    def render(row: list[str]) -> str:
        return "| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |"

    return "\n".join([
        render(normalized[0]),
        render(["---"] * width),
        *(render(row) for row in normalized[1:]),
    ])


def table_sections(
    rows: list[list[Any]],
    *,
    heading: str,
    page: int | None = None,
    location: dict[str, Any] | None = None,
) -> list[TextSection]:
    cleaned = [[_clean_cell(cell) for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return []
    header = cleaned[0]
    data_rows = cleaned[1:]
    groups: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in data_rows or [[]]:
        candidate = [header, *current, row] if data_rows else [header]
        if current and len(_render_table(candidate)) > TARGET_CHUNK_CHARS:
            groups.append([header, *current])
            current = [row]
        else:
            current.append(row)
    if current:
        groups.append([header, *current] if data_rows else [header])
    return [
        TextSection(
            text=f"表格：{heading or '未命名表格'}\n{_render_table(group)}",
            page=page,
            heading=heading,
            content_kind="table",
            location=location,
        )
        for group in groups
    ]


def _extract_pdf_sections(
    path: Path,
    ocr_engine: OcrEngine,
    on_ocr: Callable[[list[int]], None] | None,
) -> list[TextSection]:
    import pymupdf

    sections: list[TextSection] = []
    missing_pages: list[int] = []
    page_sizes: dict[int, tuple[float, float]] = {}
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            page_sizes[page_index] = (float(page.rect.width), float(page.rect.height))
            blocks = [block for block in page.get_text("blocks") if len(block) >= 7 and int(block[6]) == 0 and str(block[4]).strip()]
            page_text = "\n".join(str(block[4]).strip() for block in blocks)
            if len(re.sub(r"\s+", "", page_text)) < MIN_PDF_TEXT_CHARS:
                missing_pages.append(page_index)
                continue
            for block in sorted(blocks, key=lambda item: (float(item[1]), float(item[0]))):
                sections.append(TextSection(
                    text=str(block[4]).strip(),
                    page=page_index,
                    heading=f"第 {page_index} 页",
                    location={
                        "kind": "pdf-bbox",
                        "page": page_index,
                        "bbox": [round(float(value), 2) for value in block[:4]],
                        "page_width": round(float(page.rect.width), 2),
                        "page_height": round(float(page.rect.height), 2),
                    },
                ))
    if missing_pages:
        if on_ocr:
            on_ocr(missing_pages)
        recognized = ocr_engine.recognize_pdf_pages(path, missing_pages)
        for page_number in missing_pages:
            ocr_text = recognized.get(page_number, "").strip()
            if not ocr_text:
                continue
            width, height = page_sizes[page_number]
            sections.append(TextSection(
                text=ocr_text,
                page=page_number,
                heading=f"第 {page_number} 页（OCR）",
                content_kind="ocr",
                location={
                    "kind": "pdf-page",
                    "page": page_number,
                    "bbox": [0.0, 0.0, width, height],
                    "page_width": width,
                    "page_height": height,
                },
            ))
    return sorted(sections, key=lambda section: section.page or 0)


def _extract_docx_sections(path: Path) -> list[TextSection]:
    document = WordDocument(str(path))
    sections: list[TextSection] = []
    heading = ""
    iterator = document.iter_inner_content() if hasattr(document, "iter_inner_content") else [*document.paragraphs, *document.tables]
    for block in iterator:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            if block.style and block.style.name.lower().startswith("heading"):
                heading = text
            else:
                sections.append(TextSection(text=text, heading=heading))
        elif isinstance(block, Table):
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            sections.extend(table_sections(rows, heading=heading))
    return sections


def _extract_xlsx_sections(path: Path) -> list[TextSection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sections: list[TextSection] = []
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sections.extend(table_sections(rows, heading=f"工作表：{worksheet.title}", location={"kind": "worksheet", "sheet": worksheet.title}))
        return sections
    finally:
        workbook.close()


def _extract_xls_sections(path: Path) -> list[TextSection]:
    """Extract legacy BIFF .xls workbooks without requiring Excel/LibreOffice."""

    import xlrd

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sections: list[TextSection] = []
        for sheet_name in workbook.sheet_names():
            worksheet = workbook.sheet_by_name(sheet_name)
            rows: list[list[Any]] = []
            for row_index in range(worksheet.nrows):
                row: list[Any] = []
                for column_index in range(worksheet.ncols):
                    cell = worksheet.cell(row_index, column_index)
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
                        value: Any = ""
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat(sep=" ", timespec="seconds")
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = "是" if cell.value else "否"
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                        value = int(cell.value)
                    else:
                        value = cell.value
                    row.append(value)
                rows.append(row)
            sections.extend(table_sections(
                rows,
                heading=f"工作表：{sheet_name}",
                location={"kind": "worksheet", "sheet": sheet_name},
            ))
        return sections
    finally:
        workbook.release_resources()


def _extract_pptx_sections(path: Path) -> list[TextSection]:
    presentation = Presentation(str(path))
    sections: list[TextSection] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        heading = ""
        if slide.shapes.title is not None:
            heading = slide.shapes.title.text.strip()
        heading = heading or f"幻灯片 {slide_number}"
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                sections.extend(table_sections(rows, heading=heading, page=slide_number, location={"kind": "slide", "slide": slide_number}))
                continue
            if not getattr(shape, "has_text_frame", False) or shape is slide.shapes.title:
                continue
            text = "\n".join(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
            if text:
                sections.append(TextSection(
                    text=text,
                    page=slide_number,
                    heading=heading,
                    location={"kind": "slide", "slide": slide_number},
                ))
    return sections


def _extract_html_sections(path: Path) -> list[TextSection]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8-sig", errors="replace"), "html.parser")
    for unwanted in soup(["script", "style", "noscript", "template"]):
        unwanted.decompose()
    sections: list[TextSection] = []
    heading = ""
    for element in soup.select("h1,h2,h3,h4,h5,h6,p,li,table"):
        if element.name and element.name.startswith("h"):
            heading = element.get_text(" ", strip=True)
            continue
        if element.name == "table":
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"], recursive=False)]
                for row in element.find_all("tr")
            ]
            sections.extend(table_sections(rows, heading=heading or (soup.title.string.strip() if soup.title and soup.title.string else "HTML 表格")))
            continue
        text = element.get_text(" ", strip=True)
        if text:
            sections.append(TextSection(text=text, heading=heading))
    if not sections:
        text = soup.get_text("\n", strip=True)
        if text:
            sections.append(TextSection(text=text, heading=soup.title.string.strip() if soup.title and soup.title.string else ""))
    return sections


def extract_sections(
    path: Path,
    source_type: str,
    *,
    ocr_engine: OcrEngine | None = None,
    on_ocr: Callable[[list[int]], None] | None = None,
) -> list[TextSection]:
    if source_type == "pdf":
        return _extract_pdf_sections(path, ocr_engine or RapidOcrEngine(), on_ocr)
    if source_type == "docx":
        return _extract_docx_sections(path)
    if source_type == "xlsx":
        return _extract_xlsx_sections(path)
    if source_type == "xls":
        return _extract_xls_sections(path)
    if source_type == "pptx":
        return _extract_pptx_sections(path)
    if source_type == "html":
        return _extract_html_sections(path)
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if source_type == "md":
        sections: list[TextSection] = []
        heading = ""
        buffer: list[str] = []
        for line in text.splitlines():
            if re.match(r"^#{1,6}\s+", line):
                if buffer:
                    sections.append(TextSection(text="\n".join(buffer).strip(), heading=heading))
                    buffer = []
                heading = re.sub(r"^#{1,6}\s+", "", line).strip()
            else:
                buffer.append(line)
        if buffer:
            sections.append(TextSection(text="\n".join(buffer).strip(), heading=heading))
        return sections
    return [TextSection(text=text)]


def _split_text(text: str, max_chars: int, overlap: int = 0) -> list[str]:
    pieces: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            boundary = max(
                text.rfind("\n", start, end),
                text.rfind("。", start, end),
                text.rfind("；", start, end),
                text.rfind(". ", start, end),
            )
            if boundary > start + max_chars // 2:
                end = boundary + 1
        content = text[start:end].strip()
        if content:
            pieces.append(content)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return pieces


def _location_scope(section: TextSection) -> tuple[str, object] | None:
    location = section.location or {}
    kind = str(location.get("kind") or "")
    if kind.startswith("pdf-"):
        return "pdf", location.get("page", section.page)
    if kind == "slide":
        return "slide", location.get("slide", section.page)
    if kind == "worksheet":
        return "worksheet", location.get("sheet")
    return None


def _requires_parent_break(previous: TextSection, current: TextSection) -> bool:
    if previous.content_kind != current.content_kind:
        return True
    previous_scope = _location_scope(previous)
    current_scope = _location_scope(current)
    if previous_scope and current_scope and previous_scope[0] in {"slide", "worksheet"} and previous_scope != current_scope:
        return True
    # A parent is deliberately broader than a heading. Child chunks still stop
    # at heading/page boundaries so retrieval remains precise, while synthesis
    # can recover the surrounding clauses and neighbouring sections.
    return False


def _normalized_sections(sections: list[TextSection]) -> list[TextSection]:
    normalized: list[TextSection] = []
    for section in sections:
        text = re.sub(r"[ \t]+", " ", section.text).strip()
        if not text:
            continue
        for piece in _split_text(text, PARENT_CHUNK_CHARS):
            normalized.append(TextSection(
                text=piece,
                page=section.page,
                heading=section.heading,
                content_kind=section.content_kind,
                location=section.location,
            ))
    return normalized


def _parent_groups(sections: list[TextSection]) -> list[list[TextSection]]:
    groups: list[list[TextSection]] = []
    current: list[TextSection] = []
    current_chars = 0
    for section in _normalized_sections(sections):
        separator_chars = 1 if current else 0
        must_break = bool(current) and (
            _requires_parent_break(current[-1], section)
            or current_chars + separator_chars + len(section.text) > PARENT_CHUNK_CHARS
        )
        if must_break:
            groups.append(current)
            current = []
            current_chars = 0
            separator_chars = 0
        current.append(section)
        current_chars += separator_chars + len(section.text)
    if current:
        groups.append(current)
    return groups


def _merge_locations(sections: list[TextSection]) -> dict[str, Any] | None:
    locations = [section.location for section in sections if section.location]
    if not locations:
        return None
    first = dict(locations[0])
    if all(location == locations[0] for location in locations):
        return first
    if all(
        location.get("kind") == "pdf-bbox"
        and location.get("page") == first.get("page")
        and isinstance(location.get("bbox"), list)
        and len(location["bbox"]) == 4
        for location in locations
    ):
        boxes = [location["bbox"] for location in locations]
        first["bbox"] = [
            min(float(box[0]) for box in boxes),
            min(float(box[1]) for box in boxes),
            max(float(box[2]) for box in boxes),
            max(float(box[3]) for box in boxes),
        ]
        return first
    if all(_location_scope(section) == _location_scope(sections[0]) for section in sections):
        return first
    return None


def _child_groups(parent: list[TextSection]) -> list[list[TextSection]]:
    """Keep physical page/slide/sheet locations exact while aggregating short paragraphs."""

    groups: list[list[TextSection]] = []
    current: list[TextSection] = []
    current_scope: tuple[str, object] | None = None
    for section in parent:
        scope = _location_scope(section)
        heading_changed = bool(current) and section.heading != current[-1].heading
        scope_changed = bool(current) and scope != current_scope and (scope is not None or current_scope is not None)
        if current and (heading_changed or scope_changed):
            groups.append(current)
            current = []
        current.append(section)
        current_scope = scope
    if current:
        groups.append(current)
    return groups


def chunk_sections(sections: list[TextSection]) -> list[dict[str, Any]]:
    """Create small retrieval children backed by larger, shared structural parents."""

    chunks: list[dict[str, Any]] = []
    for parent_index, parent in enumerate(_parent_groups(sections)):
        parent_content = "\n".join(section.text for section in parent).strip()
        for child_group in _child_groups(parent):
            child_text = "\n".join(section.text for section in child_group).strip()
            location = _merge_locations(child_group)
            page = next((section.page for section in child_group if section.page is not None), None)
            heading = next((section.heading for section in child_group if section.heading), "")
            content_kind = child_group[0].content_kind
            for content in _split_text(child_text, TARGET_CHUNK_CHARS, CHUNK_OVERLAP_CHARS):
                chunks.append({
                    "content": content,
                    "page_start": page,
                    "heading": heading,
                    "content_kind": content_kind,
                    "parent_index": parent_index,
                    "parent_content": parent_content,
                    "location": location,
                })
    return chunks
