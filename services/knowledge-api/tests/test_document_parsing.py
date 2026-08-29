from __future__ import annotations

from docx import Document
import pymupdf
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import xlwt

from knowledge_api.document_parsing import RapidOcrEngine, TextSection, chunk_sections, extract_sections


def test_docx_table_is_preserved_as_table_chunk(tmp_path) -> None:
    path = tmp_path / "费用标准.docx"
    document = Document()
    document.add_heading("差旅标准", level=1)
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(
        table.rows,
        [["城市", "人员", "住宿上限"], ["上海", "普通员工", "500元"], ["上海", "部门负责人", "650元"]],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(path)

    chunks = chunk_sections(extract_sections(path, "docx"))

    assert len(chunks) == 1
    assert chunks[0]["content_kind"] == "table"
    assert chunks[0]["heading"] == "差旅标准"
    assert "| 上海 | 普通员工 | 500元 |" in chunks[0]["content"]


def test_xlsx_worksheets_are_extracted_as_table_chunks(tmp_path) -> None:
    path = tmp_path / "采购清单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "采购审批"
    worksheet.append(["金额范围", "审批人"])
    worksheet.append(["50000元以内", "部门负责人"])
    workbook.save(path)
    workbook.close()

    chunks = chunk_sections(extract_sections(path, "xlsx"))

    assert chunks[0]["content_kind"] == "table"
    assert chunks[0]["heading"] == "工作表：采购审批"
    assert "50000元以内" in chunks[0]["content"]


def test_legacy_xls_worksheets_are_extracted_as_table_chunks(tmp_path) -> None:
    path = tmp_path / "差旅标准.xls"
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("住宿标准")
    for row_index, values in enumerate([
        ["城市", "人员", "住宿上限"],
        ["上海", "普通员工", 500],
        ["上海", "部门负责人", 650],
    ]):
        for column_index, value in enumerate(values):
            worksheet.write(row_index, column_index, value)
    workbook.save(str(path))

    chunks = chunk_sections(extract_sections(path, "xls"))

    assert chunks[0]["content_kind"] == "table"
    assert chunks[0]["heading"] == "工作表：住宿标准"
    assert "| 上海 | 部门负责人 | 650 |" in chunks[0]["content"]


def test_pptx_text_and_tables_preserve_slide_location(tmp_path) -> None:
    path = tmp_path / "项目汇报.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "交付计划"
    text_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5), Inches(1))
    text_box.text_frame.text = "第一阶段在 9 月 12 日完成验收。"
    table = slide.shapes.add_table(2, 2, Inches(0.8), Inches(2.8), Inches(5), Inches(1.3)).table
    table.cell(0, 0).text = "里程碑"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "试运行"
    table.cell(1, 1).text = "李明"
    presentation.save(path)

    chunks = chunk_sections(extract_sections(path, "pptx"))

    assert any("9 月 12 日" in chunk["content"] for chunk in chunks)
    table_chunk = next(chunk for chunk in chunks if chunk["content_kind"] == "table")
    assert "| 试运行 | 李明 |" in table_chunk["content"]
    assert table_chunk["heading"] == "交付计划"
    assert table_chunk["page_start"] == 1
    assert table_chunk["location"] == {"kind": "slide", "slide": 1}


def test_html_headings_text_and_tables_are_extracted_without_scripts(tmp_path) -> None:
    path = tmp_path / "制度.html"
    path.write_text(
        """<!doctype html><html><head><title>安全制度</title><script>不要索引这段脚本</script></head>
        <body><h1>账号安全</h1><p>密码至少每 90 天更换一次。</p>
        <table><tr><th>角色</th><th>要求</th></tr><tr><td>管理员</td><td>启用双因素认证</td></tr></table></body></html>""",
        encoding="utf-8",
    )

    chunks = chunk_sections(extract_sections(path, "html"))

    assert any(chunk["heading"] == "账号安全" and "90 天" in chunk["content"] for chunk in chunks)
    table_chunk = next(chunk for chunk in chunks if chunk["content_kind"] == "table")
    assert "| 管理员 | 启用双因素认证 |" in table_chunk["content"]
    assert all("不要索引这段脚本" not in chunk["content"] for chunk in chunks)


def test_long_section_uses_shared_bounded_parent_context() -> None:
    paragraph = "项目代码 ZY-2026 的交付验收需要测试报告、部署清单和遗留问题列表。"
    # A single semantic section becomes multiple retrieval children under one bounded parent window.
    chunks = chunk_sections([TextSection(text="\n".join(paragraph for _ in range(55)), heading="交付验收")])

    assert len(chunks) >= 3
    first_parent = [chunk for chunk in chunks if chunk["parent_index"] == chunks[0]["parent_index"]]
    assert len(first_parent) >= 2
    assert len({chunk["parent_content"] for chunk in first_parent}) == 1
    assert len(first_parent[0]["parent_content"]) <= 3600
    assert len(first_parent[0]["parent_content"]) > len(first_parent[0]["content"])


def test_adjacent_short_paragraphs_share_structural_parent_context() -> None:
    sections = [
        TextSection(text=f"第 {index} 条：差旅申请应在出发前完成审批，并保留对应单据和审批记录。", heading="差旅审批")
        for index in range(1, 36)
    ]

    chunks = chunk_sections(sections)

    assert len(chunks) >= 2
    assert len({chunk["parent_index"] for chunk in chunks}) == 1
    assert len(chunks[0]["parent_content"]) > len(chunks[0]["content"])
    assert "第 1 条" in chunks[0]["parent_content"]
    assert "第 35 条" in chunks[0]["parent_content"]


def test_adjacent_headings_keep_child_labels_but_share_parent_context() -> None:
    sections = [
        TextSection(text="普通员工住宿上限为每晚 500 元。", heading="普通员工"),
        TextSection(text="部门负责人住宿上限为每晚 650 元。", heading="部门负责人"),
        TextSection(text="总经理住宿上限按公司专项审批执行。", heading="总经理"),
    ]

    chunks = chunk_sections(sections)

    assert [chunk["heading"] for chunk in chunks] == ["普通员工", "部门负责人", "总经理"]
    assert len({chunk["parent_index"] for chunk in chunks}) == 1
    assert all("500 元" in chunk["parent_content"] and "650 元" in chunk["parent_content"] for chunk in chunks)


def test_pdf_pages_keep_precise_locations_under_shared_parent() -> None:
    sections = [
        TextSection(
            text=("第一页制度内容。" * 70),
            page=1,
            heading="第 1 页",
            location={"kind": "pdf-bbox", "page": 1, "bbox": [10, 20, 300, 400], "page_width": 600, "page_height": 800},
        ),
        TextSection(
            text=("第二页审批内容。" * 70),
            page=2,
            heading="第 2 页",
            location={"kind": "pdf-bbox", "page": 2, "bbox": [15, 25, 310, 420], "page_width": 600, "page_height": 800},
        ),
    ]

    chunks = chunk_sections(sections)

    assert {chunk["page_start"] for chunk in chunks} == {1, 2}
    assert len({chunk["parent_index"] for chunk in chunks}) == 1
    assert {chunk["location"]["page"] for chunk in chunks} == {1, 2}


def test_scanned_pdf_uses_local_rapidocr(tmp_path) -> None:
    path = tmp_path / "scanned-policy.pdf"
    source = pymupdf.open()
    page = source.new_page(width=900, height=260)
    page.insert_text((50, 130), "POLICY XZ-2048 LIMIT 5000 USD", fontsize=34, fontname="helv")
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), colorspace=pymupdf.csRGB, alpha=False)
    scanned = pymupdf.open()
    scanned_page = scanned.new_page(width=900, height=260)
    scanned_page.insert_image(scanned_page.rect, stream=pixmap.tobytes("png"))
    scanned.save(path)
    scanned.close()
    source.close()

    sections = extract_sections(path, "pdf", ocr_engine=RapidOcrEngine())

    assert sections[0].content_kind == "ocr"
    assert "2048" in sections[0].text
    assert "5000" in sections[0].text
