from __future__ import annotations

from pathlib import Path
from io import BytesIO
import json

import httpx
import pytest
import pymupdf
from pptx import Presentation
from pptx.util import Inches

from knowledge_api.app import create_app

TOKEN = "test-runtime-token"
HEADERS = {"Authorization": f"Bearer {TOKEN}"}


@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"


def api_client(data_dir: Path) -> httpx.AsyncClient:
    transport = httpx.ASGITransport(app=create_app(TOKEN, data_dir=data_dir, default_model_provider="extractive"))
    return httpx.AsyncClient(transport=transport, base_url="http://knowledge-api.test")


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "上线窗口"
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6), Inches(1))
    box.text_frame.text = "生产环境上线窗口为周六 22:00。"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def make_text_pdf(text: str = "PDF POLICY ZY-7788 RETENTION 180 DAYS") -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 120), text, fontsize=14, fontname="helv")
    content = document.tobytes()
    document.close()
    return content


@pytest.mark.anyio
async def test_seeded_knowledge_bases_are_persistent(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        response = await client.get("/api/v1/knowledge-bases", headers=HEADERS)
    assert response.status_code == 200
    assert [item["id"] for item in response.json()["items"]] == ["enterprise", "project", "personal"]
    assert (tmp_path / "knowledge.db").exists()


@pytest.mark.anyio
async def test_upload_index_search_and_delete_text_document(tmp_path: Path) -> None:
    content = "# 差旅规定\n上海普通员工住宿标准为每晚500元。部门负责人住宿标准为每晚650元。"
    async with api_client(tmp_path) as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("差旅规定.md", content.encode("utf-8"), "text/markdown"))],
        )
        assert upload.status_code == 202
        document_id = upload.json()["items"][0]["document"]["id"]

        documents = await client.get("/api/v1/knowledge-bases/enterprise/documents", headers=HEADERS)
        assert documents.json()["items"][0]["status"] == "ready"
        assert documents.json()["items"][0]["chunk_count"] == 1

        jobs = await client.get("/api/v1/jobs", headers=HEADERS)
        assert jobs.json()["items"][0]["status"] == "succeeded"

        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "上海住宿标准", "knowledge_base_ids": ["enterprise"], "limit": 5},
        )
        assert search.status_code == 200
        assert search.json()["items"][0]["document_name"] == "差旅规定.md"
        assert "500元" in search.json()["items"][0]["content"]

        deleted = await client.delete(f"/api/v1/documents/{document_id}", headers=HEADERS)
        assert deleted.status_code == 204
    assert not list((tmp_path / "documents" / "enterprise").glob("*.md"))


@pytest.mark.anyio
async def test_search_can_be_limited_to_selected_documents(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        first = await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("上海制度.txt", "住宿标准：上海每晚500元。".encode(), "text/plain"))],
        )
        second = await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("北京制度.txt", "住宿标准：北京每晚450元。".encode(), "text/plain"))],
        )
        first_id = first.json()["items"][0]["document"]["id"]
        second_id = second.json()["items"][0]["document"]["id"]

        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={
                "query": "住宿标准",
                "knowledge_base_ids": ["enterprise"],
                "document_ids": [second_id],
                "limit": 5,
            },
        )

    assert search.status_code == 200
    assert search.json()["items"]
    assert {item["document_id"] for item in search.json()["items"]} == {second_id}
    assert first_id not in {item["document_id"] for item in search.json()["items"]}


@pytest.mark.anyio
async def test_rejects_document_outside_requested_knowledge_base(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/project/documents",
            headers=HEADERS,
            files=[("files", ("项目计划.txt", "项目交付计划。".encode(), "text/plain"))],
        )
        document_id = upload.json()["items"][0]["document"]["id"]
        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "项目", "knowledge_base_ids": ["enterprise"], "document_ids": [document_id]},
        )

    assert search.status_code == 400
    assert search.json()["error"]["code"] == "DOCUMENT_SCOPE_INVALID"


@pytest.mark.anyio
async def test_open_document_uses_controlled_stored_path(tmp_path: Path) -> None:
    opened: list[Path] = []
    app = create_app(
        TOKEN,
        data_dir=tmp_path,
        default_model_provider="extractive",
        document_opener=opened.append,
    )
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://knowledge-api.test") as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/personal/documents",
            headers=HEADERS,
            files=[("files", ("个人笔记.md", "# 本地笔记".encode(), "text/markdown"))],
        )
        document_id = upload.json()["items"][0]["document"]["id"]
        opened_response = await client.post(f"/api/v1/documents/{document_id}/open", headers=HEADERS)
        missing_response = await client.post("/api/v1/documents/not-found/open", headers=HEADERS)

    assert opened_response.status_code == 200
    assert opened_response.json()["status"] == "opened"
    assert len(opened) == 1
    assert opened[0].is_file()
    assert opened[0].is_relative_to((tmp_path / "documents").resolve())
    assert missing_response.status_code == 404
    assert missing_response.json()["error"]["code"] == "DOCUMENT_NOT_FOUND"


@pytest.mark.anyio
async def test_rejects_unsupported_document_type(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        response = await client.post(
            "/api/v1/knowledge-bases/personal/documents",
            headers=HEADERS,
            files=[("files", ("archive.zip", b"not-a-real-archive", "application/zip"))],
        )
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "DOCUMENT_TYPE_UNSUPPORTED"


@pytest.mark.anyio
async def test_mixed_upload_batch_is_rejected_before_any_file_is_saved(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        response = await client.post(
            "/api/v1/knowledge-bases/personal/documents",
            headers=HEADERS,
            files=[
                ("files", ("可用资料.md", b"# local policy", "text/markdown")),
                ("files", ("不支持的附件.zip", b"not-a-real-archive", "application/zip")),
            ],
        )
        documents = await client.get("/api/v1/knowledge-bases/personal/documents", headers=HEADERS)

    assert response.status_code == 415
    assert "本批次未导入任何文件" in response.json()["error"]["message"]
    assert documents.json()["items"] == []


@pytest.mark.anyio
async def test_streaming_chat_uses_real_retrieval_and_citations(tmp_path: Path) -> None:
    content = "差旅制度规定：上海普通员工住宿上限为每晚500元，部门负责人为每晚650元。"
    async with api_client(tmp_path) as client:
        await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("差旅制度.txt", content.encode("utf-8"), "text/plain"))],
        )
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={"message": "上海住宿上限是多少？", "knowledge_base_ids": ["enterprise"], "stream": True},
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/x-ndjson")
    events = [json.loads(line) for line in response.text.splitlines()]
    assert events[0]["event"] == "message.start"
    assert set(events[0]["timings_ms"]) == {"planning", "retrieval", "answer_setup", "before_stream"}
    assert all(value >= 0 for value in events[0]["timings_ms"].values())
    assert events[0]["diagnostics"]["query_context"] == "independent"
    assert events[0]["diagnostics"]["retrieval_query_count"] >= 1
    assert events[0]["diagnostics"]["retrieval_candidate_count"] >= 1
    assert events[0]["diagnostics"]["embedding_execution_mode"] == "local"
    assert any(event["event"] == "retrieval.completed" and event["count"] >= 1 for event in events)
    citation = next(event["citation"] for event in events if event["event"] == "citation")
    assert citation["document_name"] == "差旅制度.txt"
    answer = "".join(event["delta"] for event in events if event["event"] == "message.delta")
    assert "500元" in answer
    assert "[[1]]" in answer
    assert events[-1]["event"] == "message.completed"
    assert events[-1]["no_evidence"] is False
    assert events[-1]["timings_ms"]["first_answer_delta"] is not None
    assert events[-1]["timings_ms"]["total"] >= events[-1]["timings_ms"]["generation"]
    assert events[-1]["diagnostics"]["answer_path"] == "extractive-fallback"
    assert events[-1]["diagnostics"]["citation_count"] >= 1
    assert events[-1]["diagnostics"]["visible_character_count"] == len(answer)


@pytest.mark.anyio
async def test_follow_up_chat_uses_recent_history_for_retrieval(tmp_path: Path) -> None:
    content = "差旅制度规定：上海普通员工住宿上限为每晚500元，部门负责人为每晚650元。"
    async with api_client(tmp_path) as client:
        await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("差旅制度.txt", content.encode("utf-8"), "text/plain"))],
        )
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={
                "message": "那部门负责人呢？",
                "knowledge_base_ids": ["enterprise"],
                "conversation_id": "conversation-test",
                "history": [
                    {"role": "user", "content": "上海普通员工住宿上限是多少？"},
                    {"role": "assistant", "content": "普通员工每晚500元。"},
                ],
                "stream": True,
            },
        )

    answer = "".join(
        event["delta"]
        for event in (json.loads(line) for line in response.text.splitlines())
        if event["event"] == "message.delta"
    )
    assert response.status_code == 200
    assert "650元" in answer


@pytest.mark.anyio
async def test_new_short_topic_does_not_reuse_previous_allowance_answer(tmp_path: Path) -> None:
    content = (
        "# 通讯补贴\n经理通讯补贴为每月200元，总监通讯补贴为每月400元。\n\n"
        "# 考勤迟到\n迟到不足15分钟（含）的，当月第一次口头提醒，不扣款；当月第二次起每次扣款30元。"
        "迟到超过15分钟不足30分钟的，每次扣款50元。"
        "迟到超过30分钟不足2小时的，按旷工半天处理。迟到2小时及以上的，按旷工一天处理。"
    )
    async with api_client(tmp_path) as client:
        await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("人事制度.md", content.encode("utf-8"), "text/markdown"))],
        )
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={
                "message": "迟到20分钟怎么算？",
                "knowledge_base_ids": ["enterprise"],
                "history": [
                    {"role": "user", "content": "经理和总监通讯补贴差多少？"},
                    {"role": "assistant", "content": "两者相差200元。"},
                ],
                "stream": True,
            },
        )

    answer = "".join(
        event["delta"]
        for event in (json.loads(line) for line in response.text.splitlines())
        if event["event"] == "message.delta"
    )
    assert response.status_code == 200
    assert "20分钟" in answer
    assert "50元" in answer
    assert "通讯补贴" not in answer


@pytest.mark.anyio
@pytest.mark.parametrize(
    ("name", "content", "mime_type", "query", "expected"),
    [
        ("上线计划.pptx", make_pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "生产上线窗口", "周六 22:00"),
        ("账号制度.html", "<h1>账号安全</h1><p>管理员必须启用双因素认证。</p>".encode(), "text/html", "管理员认证要求", "双因素认证"),
    ],
    ids=["pptx", "html"],
)
async def test_pptx_and_html_uploads_are_searchable(
    tmp_path: Path,
    name: str,
    content: bytes,
    mime_type: str,
    query: str,
    expected: str,
) -> None:
    async with api_client(tmp_path) as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/project/documents",
            headers=HEADERS,
            files=[("files", (name, content, mime_type))],
        )
        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": query, "knowledge_base_ids": ["project"], "limit": 5},
        )

    assert upload.status_code == 202
    assert search.status_code == 200
    assert search.json()["items"]
    assert expected in search.json()["items"][0]["content"]


@pytest.mark.anyio
async def test_pdf_search_exposes_location_and_preview_returns_png(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("留存制度.pdf", make_text_pdf(), "application/pdf"))],
        )
        document_id = upload.json()["items"][0]["document"]["id"]
        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "ZY-7788 保留多久", "knowledge_base_ids": ["enterprise"], "limit": 5},
        )
        result = search.json()["items"][0]
        preview = await client.get(
            f"/api/v1/documents/{document_id}/chunks/{result['chunk_id']}/preview",
            headers=HEADERS,
        )

    assert result["location"]["kind"] == "pdf-bbox"
    assert result["location"]["page"] == 1
    assert len(result["location"]["bbox"]) == 4
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/png"
    assert preview.headers["cache-control"] == "no-store"
    assert preview.content.startswith(b"\x89PNG\r\n\x1a\n")


@pytest.mark.anyio
async def test_preview_rejects_non_pdf_and_mismatched_document_chunk(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        text_upload = await client.post(
            "/api/v1/knowledge-bases/personal/documents",
            headers=HEADERS,
            files=[("files", ("本地说明.txt", "本地说明内容".encode(), "text/plain"))],
        )
        text_document_id = text_upload.json()["items"][0]["document"]["id"]
        text_search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "本地说明", "knowledge_base_ids": ["personal"]},
        )
        text_chunk_id = text_search.json()["items"][0]["chunk_id"]
        unsupported = await client.get(
            f"/api/v1/documents/{text_document_id}/chunks/{text_chunk_id}/preview",
            headers=HEADERS,
        )

        pdf_upload = await client.post(
            "/api/v1/knowledge-bases/personal/documents",
            headers=HEADERS,
            files=[("files", ("另一份.pdf", make_text_pdf("PDF POLICY ZY-9911 ACTIVE"), "application/pdf"))],
        )
        pdf_document_id = pdf_upload.json()["items"][0]["document"]["id"]
        mismatch = await client.get(
            f"/api/v1/documents/{pdf_document_id}/chunks/{text_chunk_id}/preview",
            headers=HEADERS,
        )

    assert unsupported.status_code == 415
    assert unsupported.json()["error"]["code"] == "DOCUMENT_PREVIEW_UNSUPPORTED"
    assert mismatch.status_code == 404
    assert mismatch.json()["error"]["code"] == "DOCUMENT_LOCATION_NOT_FOUND"


@pytest.mark.anyio
async def test_chat_citation_contains_parent_context_and_location_metadata(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("定位制度.pdf", make_text_pdf("PDF POLICY ZY-7788 RETENTION 180 DAYS"), "application/pdf"))],
        )
        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "ZY-7788 RETENTION 180 DAYS", "knowledge_base_ids": ["enterprise"]},
        )
        assert search.json()["items"]
        assert search.json()["items"][0]["evidence_score"] >= 0.45
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={"message": "ZY-7788 RETENTION 180 DAYS", "knowledge_base_ids": ["enterprise"]},
        )

    events = [json.loads(line) for line in response.text.splitlines()]
    citation = next(event["citation"] for event in events if event["event"] == "citation")
    assert citation["parent_chunk_id"]
    assert citation["content_kind"] == "text"
    assert citation["location"]["kind"] == "pdf-bbox"
    assert citation["context"]
    assert "180 DAYS" in citation["context"]


@pytest.mark.anyio
async def test_runtime_reports_hybrid_index_and_reindex_metadata(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        upload = await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("错误手册.txt", "错误代码 XZ-1042 表示授权过期。".encode(), "text/plain"))],
        )
        assert upload.status_code == 202
        search = await client.post(
            "/api/v1/retrieval/search",
            headers=HEADERS,
            json={"query": "XZ-1042 是什么意思", "knowledge_base_ids": ["enterprise"]},
        )
        runtime = await client.get("/api/v1/runtime/status", headers=HEADERS)
        reindex = await client.post("/api/v1/knowledge-bases/enterprise/reindex", headers=HEADERS)

    assert search.status_code == 200
    assert search.json()["index_mode"] == "hybrid-rrf-rerank-v1"
    assert search.json()["items"][0]["retrieval_mode"] == "hybrid"
    assert search.json()["items"][0]["reranker_model_id"] == "local-feature-reranker-v1"
    assert runtime.json()["embedding"]["model_id"] == "hashing-vector-v1"
    assert runtime.json()["embedding"]["index_version"] == 2
    assert runtime.json()["reranker"]["state"] == "ready"
    assert runtime.json()["indexes"][0]["embedding_model_id"] == "hashing-vector-v1"
    assert runtime.json()["indexes"][0]["index_version"] == 2
    assert reindex.status_code == 202
    assert len(reindex.json()["items"]) == 1


@pytest.mark.anyio
async def test_streaming_chat_refuses_when_no_evidence_exists(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={"message": "不存在的产品保修期是多少？", "knowledge_base_ids": ["project"]},
        )

    events = [json.loads(line) for line in response.text.splitlines()]
    answer = "".join(event["delta"] for event in events if event["event"] == "message.delta")
    assert "没有找到足够可靠的依据" in answer
    assert not any(event["event"] == "citation" for event in events)
    assert events[-1]["no_evidence"] is True


@pytest.mark.anyio
async def test_model_settings_reject_non_local_service_url(tmp_path: Path) -> None:
    async with api_client(tmp_path) as client:
        response = await client.put(
            "/api/v1/model/settings",
            headers=HEADERS,
            json={
                "provider": "ollama",
                "base_url": "https://example.com:11434",
                "model": "qwen3:8b",
                "timeout_seconds": 90,
            },
        )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "MODEL_SETTINGS_INVALID"


@pytest.mark.anyio
async def test_ollama_provider_generates_grounded_answer_with_real_citation(tmp_path: Path) -> None:
    async def model_handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/api/tags":
            return httpx.Response(200, json={"models": [{"name": "qwen3:8b"}]})
        if request.url.path == "/api/chat":
            return httpx.Response(200, json={"message": {"content": "上海普通员工住宿上限为每晚500元。 [[1]]"}})
        return httpx.Response(404)

    app = create_app(
        TOKEN,
        data_dir=tmp_path,
        default_model_provider="auto",
        model_transport=httpx.MockTransport(model_handler),
    )
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://knowledge-api.test") as client:
        await client.post(
            "/api/v1/knowledge-bases/enterprise/documents",
            headers=HEADERS,
            files=[("files", ("差旅制度.txt", "上海普通员工住宿上限为每晚500元。".encode(), "text/plain"))],
        )
        response = await client.post(
            "/api/v1/chat/completions",
            headers=HEADERS,
            json={"message": "上海住宿上限是多少？", "knowledge_base_ids": ["enterprise"]},
        )

    events = [json.loads(line) for line in response.text.splitlines()]
    answer = "".join(event["delta"] for event in events if event["event"] == "message.delta")
    assert "500元" in answer
    assert "[[1]]" in answer
    assert events[0]["provider"] == "ollama:qwen3:8b"
