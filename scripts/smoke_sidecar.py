from __future__ import annotations

import json
import socket
import subprocess
import sys
import tempfile
import time
from io import BytesIO
from pathlib import Path

import httpx
import pymupdf
from pptx import Presentation
from pptx.util import Inches
import xlwt


ROOT = Path(__file__).resolve().parents[1]


def rust_host_target() -> str:
    result = subprocess.run(["rustc", "-vV"], check=True, capture_output=True, text=True)
    return next(line.removeprefix("host: ") for line in result.stdout.splitlines() if line.startswith("host: "))


def free_port() -> int:
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        return int(listener.getsockname()[1])


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Release window"
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6), Inches(1))
    box.text_frame.text = "Project ZY-8200 launches on Saturday at 22:00."
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def make_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 120), "PDF POLICY ZY-7788 RETENTION 180 DAYS", fontsize=14, fontname="helv")
    output = document.tobytes()
    document.close()
    return output


def make_xls() -> bytes:
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Legacy policy")
    worksheet.write(0, 0, "Policy code")
    worksheet.write(0, 1, "Approval threshold")
    worksheet.write(1, 0, "ZY-8400")
    worksheet.write(1, 1, 7650)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def wait_until_ready(client: httpx.Client, headers: dict[str, str]) -> None:
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        try:
            response = client.get("/health", headers=headers)
            if response.is_success:
                return
        except httpx.HTTPError:
            pass
        time.sleep(0.25)
    raise RuntimeError("Sidecar did not become ready within 180 seconds")


def upload(client: httpx.Client, headers: dict[str, str], name: str, content: bytes, mime_type: str) -> str:
    response = client.post(
        "/api/v1/knowledge-bases/enterprise/documents",
        headers=headers,
        files={"files": (name, content, mime_type)},
    )
    response.raise_for_status()
    return str(response.json()["items"][0]["document"]["id"])


def wait_for_documents_ready(client: httpx.Client, headers: dict[str, str], expected_count: int) -> None:
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        response = client.get("/api/v1/knowledge-bases/enterprise/documents", headers=headers)
        response.raise_for_status()
        documents = response.json()["items"]
        failed = [item for item in documents if item["status"] == "failed"]
        if failed:
            raise AssertionError(f"Sidecar indexing failed: {failed[0].get('error_message') or failed[0]['source_name']}")
        if len(documents) >= expected_count and all(item["status"] == "ready" for item in documents):
            return
        time.sleep(0.25)
    raise RuntimeError("Documents did not finish indexing within 180 seconds")


def search(client: httpx.Client, headers: dict[str, str], query: str) -> dict[str, object]:
    response = client.post(
        "/api/v1/retrieval/search",
        headers=headers,
        json={"query": query, "knowledge_base_ids": ["enterprise"], "limit": 5},
    )
    response.raise_for_status()
    items = response.json()["items"]
    if not items:
        raise AssertionError(f"No result returned for {query}")
    return items[0]


def main() -> None:
    target = rust_host_target()
    executable = ROOT / "src-tauri" / "binaries" / f"knowledge-api-{target}{'.exe' if sys.platform == 'win32' else ''}"
    if not executable.is_file():
        raise SystemExit(f"Sidecar not found: {executable}. Run npm run sidecar:build first.")

    port = free_port()
    token = "zhiyu-sidecar-smoke-token"
    base_url = f"http://127.0.0.1:{port}"
    creation_flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    with tempfile.TemporaryDirectory(prefix="zhiyu-sidecar-smoke-") as data_dir:
        # Keep the packaging smoke test deterministic on clean CI runners.
        # Product defaults still use auto-detected Ollama/qwen3:8b and bge-m3;
        # this isolated temporary profile only validates the bundled local API.
        Path(data_dir, "model-settings.json").write_text(
            json.dumps({
                "provider": "extractive",
                "base_url": "http://127.0.0.1:11434",
                "model": "qwen3:8b",
                "timeout_seconds": 90,
            }, ensure_ascii=False),
            encoding="utf-8",
        )
        process = subprocess.Popen(
            [
                str(executable),
                "--port", str(port),
                "--token", token,
                "--data-dir", data_dir,
                "--embedding-provider", "hashing",
            ],
            cwd=ROOT,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=creation_flags,
        )
        headers = {"Authorization": f"Bearer {token}"}
        try:
            with httpx.Client(base_url=base_url, timeout=180, trust_env=False) as client:
                wait_until_ready(client, headers)
                runtime = client.get("/api/v1/runtime/status", headers=headers)
                runtime.raise_for_status()
                runtime_data = runtime.json()
                assert runtime_data["model"]["provider"] == "extractive"
                assert runtime_data["model"]["state"] == "extractive"
                assert runtime_data["embedding"]["model_id"] == "hashing-vector-v1"
                assert runtime_data["embedding"]["state"] == "ready"
                assert runtime_data["ocr"]["state"] == "ready"

                upload(client, headers, "release-plan.pptx", make_pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")
                upload(client, headers, "account-policy.html", "<h1>Account security</h1><p>Policy ZY-8300 requires two-factor authentication.</p>".encode(), "text/html")
                upload(client, headers, "legacy-policy.xls", make_xls(), "application/vnd.ms-excel")
                pdf_document_id = upload(client, headers, "retention-policy.pdf", make_pdf(), "application/pdf")
                wait_for_documents_ready(client, headers, 4)

                pptx_result = search(client, headers, "ZY-8200 Saturday 22:00")
                html_result = search(client, headers, "ZY-8300 two-factor authentication")
                xls_result = search(client, headers, "ZY-8400 approval threshold 7650")
                pdf_result = search(client, headers, "ZY-7788 RETENTION 180 DAYS")
                assert pptx_result["format"] == "PPTX"
                assert html_result["format"] == "HTML"
                assert xls_result["format"] == "XLS"
                assert pdf_result["location"]["kind"] == "pdf-bbox"  # type: ignore[index]
                assert pdf_result["parent_chunk_id"]
                assert pdf_result["parent_content"]

                preview = client.get(
                    f"/api/v1/documents/{pdf_document_id}/chunks/{pdf_result['chunk_id']}/preview",
                    headers=headers,
                )
                preview.raise_for_status()
                assert preview.headers["content-type"] == "image/png"
                assert preview.content.startswith(b"\x89PNG\r\n\x1a\n")

                chat = client.post(
                    "/api/v1/chat/completions",
                    headers=headers,
                    json={"message": "What is the retention period for ZY-7788?", "knowledge_base_ids": ["enterprise"]},
                )
                chat.raise_for_status()
                events = [json.loads(line) for line in chat.text.splitlines()]
                start = next(event for event in events if event["event"] == "message.start")
                assert start["provider"] == "extractive-local-v1"
                assert any(event["event"] == "citation" for event in events)
                assert events[-1]["event"] == "message.completed"

                print(json.dumps({
                    "sidecar": str(executable),
                    "model": runtime_data["model"],
                    "embedding": runtime_data["embedding"],
                    "ocr": runtime_data["ocr"],
                    "formats": [pptx_result["format"], html_result["format"], xls_result["format"], pdf_result["format"]],
                    "preview": preview.headers["content-type"],
                    "chat_provider": start["provider"],
                }, ensure_ascii=False, indent=2))
                client.post("/api/v1/runtime/shutdown", headers=headers)
        finally:
            # Request graceful shutdown on both success and assertion failure.
            # A one-file PyInstaller executable may hand work to its child
            # process, so waiting on only the original Popen handle is not a
            # sufficient cleanup guarantee on Windows.
            try:
                with httpx.Client(base_url=base_url, timeout=5, trust_env=False) as cleanup_client:
                    cleanup_client.post("/api/v1/runtime/shutdown", headers=headers)
            except httpx.HTTPError:
                pass
            try:
                process.wait(timeout=12)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)


if __name__ == "__main__":
    main()
